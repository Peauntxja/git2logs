#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
离线回归测试（无 GitLab 网络）。

覆盖：报告 golden、HTML 解析、CLI 映射、fetch_commits mock。
"""
from __future__ import annotations

import json
import tempfile
import threading
from types import SimpleNamespace
import unittest
from pathlib import Path
from unittest import mock

from tests._common import AUTHOR, EXPECTED, FIXTURES, SINCE, UNTIL
from tests.helpers.fake_gitlab import build_grouped_commits, build_sample_all_results
from tests.helpers.normalize import normalize_report_text

from cli import build_argument_parser, cli_output_format, cli_report_type_name
from excel_exporter import fill_excel_template, merge_and_normalize_tasks
from models import GitLabConnectionError, OperationCancelled, ReportGenerationError, ReportParams
from report_generator import (
    generate_daily_report,
    generate_markdown_log,
    generate_statistics_report,
    generate_work_hours_report,
)
from report_html import generate_html_report, parse_daily_report
from service import Git2LogsService
from work_hours import calculate_work_hours
from gitlab_client import (
    get_all_branches,
    get_commits_by_author,
    retry_gitlab_call,
    scan_all_projects,
)
from commit_analysis import clear_commit_cache, get_commit_details
from gui.service_bridge import ServiceBridgeMixin


def _load_expected(name: str) -> str:
    return (EXPECTED / name).read_text(encoding="utf-8")


def _fake_commit_details(_project, commit):
    msg = (commit.message or "").split("\n")[0]
    return {
        "short_message": msg,
        "full_message": commit.message or "",
        "stats": None,
        "changed_files": [],
    }


class GoldenReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not (EXPECTED / "statistics.md").exists():
            raise unittest.SkipTest(
                "缺少 tests/expected，请先运行: python3 scripts/update_golden_fixtures.py"
            )

    def test_statistics_report_matches_golden(self):
        actual = generate_statistics_report(
            build_sample_all_results(), AUTHOR,
            since_date=SINCE, until_date=UNTIL,
        )
        self.assertEqual(normalize_report_text(actual), _load_expected("statistics.md"))

    def test_daily_report_matches_golden(self):
        with mock.patch("report_generator.get_commit_details", side_effect=_fake_commit_details):
            actual = generate_daily_report(
                build_sample_all_results(), AUTHOR,
                since_date=SINCE, until_date=UNTIL, branch="main",
            )
        self.assertEqual(normalize_report_text(actual), _load_expected("daily_report.md"))

    def test_work_hours_report_matches_golden(self):
        actual = generate_work_hours_report(
            build_sample_all_results(), AUTHOR,
            since_date=SINCE, until_date=UNTIL, branch="main",
        )
        self.assertEqual(normalize_report_text(actual), _load_expected("work_hours.md"))

    def test_commits_markdown_matches_golden(self):
        actual = generate_markdown_log(
            build_grouped_commits(), AUTHOR,
            repo_name="demo/project-a", project=None,
        )
        self.assertEqual(normalize_report_text(actual), _load_expected("commits.md"))

    def test_work_hours_data_matches_golden(self):
        actual = calculate_work_hours(
            build_sample_all_results(),
            since_date=SINCE, until_date=UNTIL, branch="main",
        )
        self.assertEqual(actual, json.loads(_load_expected("work_hours_data.json")))

    def test_excel_merge_rules_match_golden(self):
        wh_data = json.loads(_load_expected("work_hours_data.json"))
        tasks = []
        for date_str, date_data in wh_data.items():
            for _proj, pdata in date_data.get("projects", {}).items():
                for task in pdata.get("tasks", []):
                    tasks.append({
                        "task_name": task["task_name"],
                        "hours": task["hours"],
                        "start_date": date_str,
                        "end_date": date_str,
                        "description": task["task_name"],
                        "task_type": task.get("task_type", ""),
                    })
        self.assertEqual(
            merge_and_normalize_tasks(tasks),
            json.loads(_load_expected("excel_merged_tasks.json")),
        )

    def test_fill_excel_template_row_count(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl 未安装")

        template = FIXTURES / "work_hours_template.xlsx"
        if not template.exists():
            self.skipTest("缺少 tests/fixtures/work_hours_template.xlsx")

        wh_data = json.loads(_load_expected("work_hours_data.json"))
        meta = json.loads(_load_expected("meta.json"))
        out = FIXTURES / "_test_export_output.xlsx"
        if out.exists():
            out.unlink()

        # 在模板副本中塞入“上月”残留行 + 非变更列，确认导出会清空旧任务但保留参考字段
        dirty = FIXTURES / "_test_dirty_template.xlsx"
        wb = openpyxl.load_workbook(template)
        ws = wb.active
        from excel_exporter import _find_header_row
        header_row = _find_header_row(ws)
        self.assertIsNotNone(header_row)
        stale_marker = "__STALE_OLD_MONTH__"
        ws.cell(row=header_row + 1, column=1, value=stale_marker)
        # 额外列：指派人/确认人/状态（非导出覆盖列）
        assignee_col = ws.max_column + 1
        ws.cell(row=header_row, column=assignee_col, value="指派任务人")
        ws.cell(row=header_row + 1, column=assignee_col, value="")
        ws.cell(row=header_row + 2, column=1, value="上月任务A")
        ws.cell(row=header_row + 2, column=assignee_col, value="张三")
        ws.cell(row=header_row + 3, column=1, value="上月任务B")
        ws.cell(row=header_row + 3, column=assignee_col, value="张三")
        wb.save(dirty)
        wb.close()

        count = fill_excel_template(
            template_path=dirty,
            work_hours_data=wh_data,
            output_path=out,
        )
        self.assertEqual(count, meta["excel_expected_row_count"])
        self.assertTrue(out.exists())

        out_wb = openpyxl.load_workbook(out)
        out_ws = out_wb.active
        values = [
            str(c.value) if c.value is not None else ""
            for row in out_ws.iter_rows()
            for c in row
        ]
        self.assertNotIn(stale_marker, values)
        self.assertNotIn("上月任务A", values)
        self.assertEqual(out_ws.max_row, header_row + count)
        for row_idx in range(header_row + 1, header_row + count + 1):
            self.assertEqual(out_ws.cell(row=row_idx, column=assignee_col).value, "张三")
        out_wb.close()
        out.unlink()
        dirty.unlink()


class ReportHtmlTests(unittest.TestCase):
    def test_parse_daily_report_matches_golden(self):
        golden_path = EXPECTED / "report_html_parsed.json"
        if not golden_path.exists():
            self.skipTest("缺少 report_html_parsed.json")

        actual = parse_daily_report(str(EXPECTED / "daily_report.md"))
        expected = json.loads(golden_path.read_text(encoding="utf-8"))
        self.assertEqual(actual, expected)

    def test_generate_html_report_includes_sorted_timeline(self):
        golden_path = EXPECTED / "report_html_parsed.json"
        if not golden_path.exists():
            self.skipTest("缺少 report_html_parsed.json")

        data = json.loads(golden_path.read_text(encoding="utf-8"))
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "daily.html"
            generate_html_report(data, str(out))
            html = out.read_text(encoding="utf-8")
        self.assertIn("10:30", html)
        self.assertIn("16:45", html)
        self.assertLess(html.index("10:30"), html.index("16:45"))

    def test_service_generates_html_daily_report(self):
        daily_report = _load_expected("daily_report.md")
        with tempfile.TemporaryDirectory() as tmp:
            params = ReportParams(
                gitlab_url="http://gitlab.example.com",
                token="token",
                author=AUTHOR,
                output_format="html",
                output_path=tmp,
                since_date=SINCE,
                until_date=UNTIL,
            )
            with mock.patch("service.generate_daily_report", return_value=daily_report):
                result = Git2LogsService()._build_report({}, params)
            output = Path(result["output_file"])
            self.assertEqual(output.name, f"{SINCE}_daily_report.html")
            self.assertTrue(output.exists())
            self.assertFalse((Path(f"{output}.source.md")).exists())

    def test_service_cancel_guard(self):
        import threading

        cancel_event = threading.Event()
        cancel_event.set()
        with self.assertRaises(OperationCancelled):
            Git2LogsService._raise_if_cancelled(cancel_event)

    def test_service_reports_png_renderer_failure(self):
        daily_report = _load_expected("daily_report.md")
        with tempfile.TemporaryDirectory() as tmp:
            params = ReportParams(
                gitlab_url="http://gitlab.example.com",
                token="token",
                author=AUTHOR,
                output_format="png",
                output_path=tmp,
                since_date=SINCE,
                until_date=UNTIL,
            )
            with (
                mock.patch("service.generate_daily_report", return_value=daily_report),
                mock.patch("image_converter.convert_html_to_image", return_value=False),
                self.assertRaises(ReportGenerationError),
            ):
                Git2LogsService()._build_report({}, params)


class CliParserTests(unittest.TestCase):
    def test_scan_all_daily_report_format_mapping(self):
        args = build_argument_parser().parse_args([
            "--scan-all", "--gitlab-url", "http://gitlab.example.com",
            "--author", "MIZUKI", "--daily-report",
        ])
        self.assertEqual(cli_output_format(args), "daily_report")
        self.assertEqual(cli_report_type_name(args), "daily_report")

    def test_single_repo_defaults_to_commits(self):
        args = build_argument_parser().parse_args([
            "--repo", "group/project", "--author", "MIZUKI",
        ])
        self.assertEqual(cli_output_format(args), "commits")
        self.assertEqual(cli_report_type_name(args), "all_projects")


class FetchCommitsTests(unittest.TestCase):
    def _params(self, **kwargs) -> ReportParams:
        base = dict(
            gitlab_url="http://gitlab.example.com",
            token="token",
            author="MIZUKI",
            scan_all=False,
            repo_url="http://gitlab.example.com/group/project.git",
        )
        base.update(kwargs)
        return ReportParams(**base)

    @mock.patch("service.create_gitlab_client")
    @mock.patch("service.get_commits_by_author")
    def test_fetch_commits_single_project_returns_results(
        self, mock_get_commits, mock_create_client,
    ):
        project = mock.Mock(name="Demo")
        commit = mock.Mock(message="feat: test")
        mock_gl = mock.Mock()
        mock_gl.projects.get.return_value = project
        mock_get_commits.return_value = [commit]
        mock_create_client.return_value = mock_gl

        result = Git2LogsService().fetch_commits(self._params())
        entry = next(iter(result.values()))
        self.assertEqual(entry["commits"], [commit])

    @mock.patch("service.create_gitlab_client")
    def test_fetch_commits_strict_raises_on_project_error(self, mock_create_client):
        mock_gl = mock.Mock()
        mock_gl.projects.get.side_effect = RuntimeError("404 Project Not Found")
        mock_create_client.return_value = mock_gl

        with self.assertRaises(GitLabConnectionError):
            Git2LogsService().fetch_commits(
                self._params(), strict_single_project=True,
            )


class ResilienceTests(unittest.TestCase):
    def test_background_thread_gets_commit_details(self):
        clear_commit_cache()
        commit = SimpleNamespace(
            id="abcdef123456",
            message="feat: test",
            committed_date="2026-01-01T00:00:00Z",
            author_name="MIZUKI",
        )
        detailed_commit = SimpleNamespace(
            stats={"additions": 2, "deletions": 1, "total": 3},
            diff=lambda: [SimpleNamespace(new_path="demo.py", old_path="", diff="+line")],
        )
        project = SimpleNamespace(id=1, commits=SimpleNamespace(get=lambda _id: detailed_commit))
        result = {}

        thread = threading.Thread(
            target=lambda: result.update(get_commit_details(project, commit)),
        )
        thread.start()
        thread.join()

        self.assertEqual(result["stats"]["total"], 3)
        self.assertEqual(result["changed_files"][0]["path"], "demo.py")

    def test_get_all_branches_paginates_beyond_one_hundred(self):
        first_page = [mock.Mock(name=f"branch-{index}") for index in range(100)]
        final_branch = mock.Mock(name="branch-100")
        project = mock.Mock()
        project.branches.list.side_effect = [first_page, [final_branch]]

        branches = get_all_branches(project)

        self.assertEqual(len(branches), 101)
        self.assertEqual(project.branches.list.call_args_list[0].kwargs["page"], 1)
        self.assertEqual(project.branches.list.call_args_list[1].kwargs["page"], 2)

    def test_retryable_request_retries_then_succeeds(self):
        response_error = RuntimeError("503 service unavailable")
        operation = mock.Mock(side_effect=[response_error, "ok"])
        with mock.patch("gitlab_client.time.sleep") as sleep:
            self.assertEqual(retry_gitlab_call(operation, "测试请求"), "ok")
        self.assertEqual(operation.call_count, 2)
        sleep.assert_called_once()

    def test_auth_error_does_not_retry(self):
        error = RuntimeError("401 unauthorized")
        operation = mock.Mock(side_effect=error)
        with mock.patch("gitlab_client.time.sleep") as sleep:
            with self.assertRaises(RuntimeError):
                retry_gitlab_call(operation, "测试请求")
        self.assertEqual(operation.call_count, 1)
        sleep.assert_not_called()

    def test_rate_limit_uses_retry_after_header(self):
        error = RuntimeError("too many requests")
        error.response_code = 429
        error.headers = {"Retry-After": "1.5"}
        operation = mock.Mock(side_effect=[error, "ok"])
        with (
            mock.patch("gitlab_client._rate_limit_until", 0),
            mock.patch("gitlab_client.time.sleep") as sleep,
        ):
            self.assertEqual(retry_gitlab_call(operation, "测试请求"), "ok")
        sleep.assert_called_once_with(1.5)

    def test_shared_rate_limit_waits_before_request(self):
        with (
            mock.patch("gitlab_client._rate_limit_until", 12),
            mock.patch("gitlab_client.time.monotonic", return_value=10),
            mock.patch("gitlab_client.time.sleep") as sleep,
        ):
            self.assertEqual(retry_gitlab_call(lambda: "ok", "测试请求"), "ok")
        sleep.assert_called_once_with(2)

    def test_all_branch_scan_deduplicates_as_it_collects(self):
        shared = SimpleNamespace(id="shared")
        first = SimpleNamespace(id="first")
        second = SimpleNamespace(id="second")
        project = mock.Mock()
        project.branches.list.return_value = [
            SimpleNamespace(name="main"),
            SimpleNamespace(name="develop"),
        ]
        project.commits.list.side_effect = lambda **params: {
            "main": [shared, first],
            "develop": [shared, second],
        }[params["ref_name"]]

        commits = get_commits_by_author(project, AUTHOR)

        self.assertEqual([commit.id for commit in commits], ["shared", "first", "second"])

    @mock.patch("gitlab_client.get_commits_by_author")
    @mock.patch("gitlab_client.get_all_projects")
    def test_scan_summary_keeps_partial_failures(
        self,
        mock_get_all_projects,
        mock_get_commits,
    ):
        first = mock.Mock(path_with_namespace="group/first")
        second = mock.Mock(path_with_namespace="group/second")
        mock_get_all_projects.return_value = [first, second]
        mock_get_commits.side_effect = [RuntimeError("503 service unavailable"), []]
        summary = {}

        result = scan_all_projects(mock.Mock(), AUTHOR, scan_summary=summary)

        self.assertEqual(result, {})
        self.assertEqual(summary["failed"], 1)
        self.assertEqual(summary["empty"], 1)

    @mock.patch("gitlab_client.get_all_projects")
    def test_scan_summary_collects_branch_metrics(self, mock_get_all_projects):
        project = mock.Mock(path_with_namespace="group/project")
        mock_get_all_projects.return_value = [project]
        summary = {}

        def get_commits(*_args, scan_metrics=None, **_kwargs):
            scan_metrics.update({"branches_scanned": 3, "branches_skipped": 1})
            return []

        with mock.patch("gitlab_client.get_commits_by_author", side_effect=get_commits):
            scan_all_projects(mock.Mock(), AUTHOR, scan_summary=summary)

        self.assertEqual(summary["branches_scanned"], 3)
        self.assertEqual(summary["branches_skipped"], 1)

    @mock.patch("gitlab_client.get_all_projects")
    def test_cancelled_scan_does_not_submit_later_projects(self, mock_get_all_projects):
        projects = [mock.Mock(path_with_namespace=f"group/project-{index}") for index in range(3)]
        mock_get_all_projects.return_value = projects
        cancel_event = threading.Event()
        scanned_projects = []

        def get_commits(project, *_args, **_kwargs):
            scanned_projects.append(project.path_with_namespace)
            cancel_event.set()
            return []

        with mock.patch("gitlab_client.get_commits_by_author", side_effect=get_commits):
            summary = {}
            scan_all_projects(
                mock.Mock(),
                AUTHOR,
                max_workers=1,
                cancel_event=cancel_event,
                scan_summary=summary,
            )

        self.assertEqual(scanned_projects, ["group/project-0"])
        self.assertTrue(summary["cancelled"])

    def test_atomic_write_preserves_existing_file_on_replace_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "report.md"
            target.write_text("old", encoding="utf-8")
            with mock.patch("service.os.replace", side_effect=OSError("disk error")):
                with self.assertRaises(OSError):
                    Git2LogsService._write_file(target, "new")
            self.assertEqual(target.read_text(encoding="utf-8"), "old")


class GuiParameterTests(unittest.TestCase):
    def test_saved_none_branch_does_not_become_branch_name(self):
        bridge = ServiceBridgeMixin()
        bridge.log = mock.Mock()
        bridge._resolve_dates_from_cached = mock.Mock(return_value=(SINCE, UNTIL))
        params = {
            "gitlab_url": "http://gitlab.example.com",
            "token": "token",
            "author": AUTHOR,
            "repo": "",
            "branch": "None",
            "output_path": "",
            "scan_all": True,
            "output_format": "daily_report",
        }

        report_params = bridge._build_report_params(params)

        self.assertIsNone(report_params.branch)


class CommitAnalysisTests(unittest.TestCase):
    def test_is_merge_commit(self):
        from commit_analysis import is_merge_commit

        self.assertTrue(is_merge_commit("Merge branch 'test' of http://gitlab.example.com/a into test"))
        self.assertTrue(is_merge_commit("merge pull request !123 from feature/foo"))
        self.assertFalse(is_merge_commit("feat(auth): 增加登录校验"))

    def test_analyze_commit_type_perf(self):
        from commit_analysis import analyze_commit_type

        self.assertEqual(analyze_commit_type("perf(views): 优化委员会组织树加载体验"), ("性能优化", "⚡"))

    def test_feat_with_fix_keyword_is_bugfix(self):
        from commit_analysis import analyze_commit_type

        self.assertEqual(analyze_commit_type("feat:修复新增时 id 传参"), ("Bug修复", "🐛"))
        self.assertEqual(analyze_commit_type("feat(auth): 增加登录校验"), ("功能开发", "✨"))

    def test_format_date_chinese_strips_whitespace(self):
        from utils.date_utils import format_date_chinese

        self.assertEqual(format_date_chinese(" 2026-06-08 "), "2026年06月08日")
        self.assertEqual(format_date_chinese("2026-08-06"), "2026年08月06日")


class BranchSkipTests(unittest.TestCase):
    def _branch(self, committed_date: str):
        commit = mock.Mock()
        commit.committed_date = committed_date
        branch = mock.Mock()
        branch.commit = commit
        branch.name = "main"
        return branch

    def test_tip_after_until_must_not_skip(self):
        """查昨天时 tip 已到今天，仍应查询（旧逻辑会误跳过）。"""
        from gitlab_client import _should_skip_branch

        branch = self._branch("2026-08-07T10:00:00Z")
        self.assertFalse(
            _should_skip_branch(branch, since_date="2026-08-06", until_date="2026-08-06")
        )

    def test_tip_before_since_skips(self):
        from gitlab_client import _should_skip_branch

        branch = self._branch("2026-08-01T10:00:00Z")
        self.assertTrue(
            _should_skip_branch(branch, since_date="2026-08-06", until_date="2026-08-06")
        )

    def test_morning_shanghai_tip_not_skipped(self):
        """上海早上 7:30 提交（UTC 前一天 23:30），查当天不应跳过。"""
        from gitlab_client import _should_skip_branch

        branch = self._branch("2026-08-05T23:30:00Z")
        self.assertFalse(
            _should_skip_branch(branch, since_date="2026-08-06", until_date="2026-08-06")
        )


class TimezoneTests(unittest.TestCase):
    def test_to_gitlab_datetime_shanghai_day(self):
        from utils.date_utils import to_gitlab_datetime

        self.assertEqual(to_gitlab_datetime("2026-08-06"), "2026-08-05T16:00:00Z")
        self.assertEqual(
            to_gitlab_datetime("2026-08-06", end_of_day=True),
            "2026-08-06T15:59:59Z",
        )

    def test_morning_commit_maps_to_shanghai_day(self):
        from utils.date_utils import to_local_date_str

        self.assertEqual(to_local_date_str("2026-08-05T23:30:00Z"), "2026-08-06")


if __name__ == "__main__":
    unittest.main()
